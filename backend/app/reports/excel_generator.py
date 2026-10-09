import io
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_report_excel(data: dict) -> io.BytesIO:
    wb = Workbook()
    
    # Styles
    title_font = Font(name="Arial", size=16, bold=True, color="1E3A8A")
    subtitle_font = Font(name="Arial", size=10, italic=True, color="64748B")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    totals_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    totals_font = Font(name="Arial", size=11, bold=True)
    
    thin_side = Side(style="thin", color="E2E8F0")
    thin_border = Border(top=thin_side, left=thin_side, right=thin_side, bottom=thin_side)
    double_bottom_border = Border(
        top=Side(style="thin", color="94A3B8"),
        bottom=Side(style="double", color="0F172A")
    )

    # -------------------------------------------------------------
    # SHEET 1: Monthly Summary
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Monthly Summary"
    ws1.freeze_panes = "A5"

    ws1.merge_cells("A1:E1")
    ws1["A1"] = "MONTHLY BREAKFAST & FINANCIAL SUMMARY"
    ws1["A1"].font = title_font

    ws1.merge_cells("A2:E2")
    ws1["A2"] = f"Year: {data.get('selectedYear', 'All Years')} | Department: {data.get('selectedDepartment', 'ALL')}"
    ws1["A2"].font = subtitle_font

    headers1 = ["Month", "Opening Balance", "Money Received", "Total Spent", "Closing Balance"]
    ws1.append([])
    ws1.append(headers1)
    
    header_row_1 = ws1[4]
    for cell in header_row_1:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center")

    for row in data.get("monthlySummary", []):
        r = [
            row.get("monthName"),
            row.get("openingBalance", 0),
            row.get("moneyReceived", 0),
            row.get("totalSpent", 0),
            row.get("closingBalance", 0)
        ]
        ws1.append(r)
        curr_row = ws1[ws1.max_row]
        curr_row[0].alignment = Alignment(vertical="center", horizontal="left")
        for c in curr_row[1:]:
            c.number_format = '"₹"#,##0.00'
            c.alignment = Alignment(vertical="center", horizontal="right")
        for c in curr_row:
            c.border = thin_border

    yearly_total = data.get("yearlyTotal")
    if yearly_total:
        ws1.append([
            "TOTAL / PERIOD SUMMARY",
            "",
            yearly_total.get("totalMoneyReceived", 0),
            yearly_total.get("totalSpent", 0),
            yearly_total.get("closingBalance", 0)
        ])
        tot_row = ws1[ws1.max_row]
        for c in tot_row:
            c.fill = totals_fill
            c.font = totals_font
            c.border = double_bottom_border
        tot_row[0].alignment = Alignment(vertical="center", horizontal="left")
        for c in tot_row[2:]:
            c.number_format = '"₹"#,##0.00'
            c.alignment = Alignment(vertical="center", horizontal="right")

    _autofit(ws1)

    # -------------------------------------------------------------
    # SHEET 2: Employee Monthly Report
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Employee Monthly Report")
    ws2.freeze_panes = "A5"

    ws2.merge_cells("A1:G1")
    ws2["A1"] = "EMPLOYEE BREAKFAST PARTICIPATION REPORT"
    ws2["A1"].font = title_font

    ws2.merge_cells("A2:G2")
    ws2["A2"] = f"Period: {data.get('selectedMonth', data.get('selectedYear', ''))} | Department: {data.get('selectedDepartment', 'ALL')}"
    ws2["A2"].font = subtitle_font

    headers2 = ["Emp ID", "Name", "Department", "Applicable Days", "Taken", "Not Taken", "No Response", "Reasons"]
    ws2.append([])
    ws2.append(headers2)

    for cell in ws2[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center")

    for emp in data.get("employeeReport", []):
        reasons_dict = emp.get("reasonBreakdown", {})
        reasons_str = ", ".join(f"{k}: {v}" for k, v in reasons_dict.items()) if reasons_dict else "-"
        ws2.append([
            emp.get("employeeId"),
            emp.get("name"),
            emp.get("department"),
            emp.get("totalDays", 0),
            emp.get("takenCount", 0),
            emp.get("notTakenCount", 0),
            emp.get("noResponseCount", 0),
            reasons_str
        ])
        for cell in ws2[ws2.max_row]:
            cell.border = thin_border

    _autofit(ws2)

    # -------------------------------------------------------------
    # SHEET 3: Order Summary
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Order Summary")
    ws3.freeze_panes = "A5"

    ws3.merge_cells("A1:H1")
    ws3["A1"] = "BREAKFAST ORDER SUMMARY"
    ws3["A1"].font = title_font

    headers3 = ["Order ID", "Date", "Order Type", "Order Title", "Time", "Headcount", "Total Qty (Actual)", "Items", "Total Cost"]
    ws3.append([])
    ws3.append([])
    ws3.append(headers3)

    for cell in ws3[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center")

    for order in data.get("orderSummary", []):
        items_str = f"BF: {order.get('breakfastItems', '')} | Common: {order.get('commonItems', '')}"
        tot_qty = order.get("totalQuantity")
        if tot_qty is None:
            tot_qty = order.get("actualResponseQuantity", order.get("applicableCount", 0))
        title = order.get("orderTitle", "")
        if order.get("clientName"):
            title = f"{title} [Client: {order.get('clientName')}]"
        ws3.append([
            order.get("orderId"),
            order.get("businessDate"),
            order.get("orderType"),
            title,
            order.get("orderTime"),
            order.get("applicableCount", 0),
            tot_qty,
            items_str,
            order.get("totalCost", 0)
        ])
        curr_row = ws3[ws3.max_row]
        curr_row[8].number_format = '"₹"#,##0.00'
        for cell in curr_row:
            cell.border = thin_border

    _autofit(ws3)

    # -------------------------------------------------------------
    # SHEET 4: Money Transactions
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Money Transactions")
    ws4.freeze_panes = "A5"

    ws4.merge_cells("A1:I1")
    ws4["A1"] = "FINANCIAL LEDGER TRANSACTIONS"
    ws4["A1"].font = title_font

    headers4 = ["Date", "Time", "Transaction ID", "Type", "Deposit", "Debited", "Balance After", "Source", "Description"]
    ws4.append([])
    ws4.append([])
    ws4.append(headers4)

    for cell in ws4[4]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", horizontal="center")

    txns = data.get("moneyTransactions", [])
    for txn in txns:
        txn_type = (txn.get("type") or "").upper()
        amt = float(txn.get("amount", 0) or 0)
        is_deposit = txn_type in ["MONEY_RECEIVED", "DEPOSIT", "CREDIT", "REVERSAL"]

        ws4.append([
            txn.get("transactionDate") or txn.get("transaction_date"),
            txn.get("transactionTime") or txn.get("transaction_time"),
            txn.get("transactionId") or txn.get("transaction_id"),
            txn.get("type"),
            amt if is_deposit else "-",
            amt if not is_deposit else "-",
            txn.get("balanceAfterTransaction") or txn.get("balance_after_transaction", 0),
            txn.get("source"),
            txn.get("description")
        ])
        curr_row = ws4[ws4.max_row]
        if is_deposit:
            curr_row[4].number_format = '"₹"#,##0.00'
            curr_row[4].alignment = Alignment(vertical="center", horizontal="right")
            curr_row[5].alignment = Alignment(vertical="center", horizontal="center")
        else:
            curr_row[4].alignment = Alignment(vertical="center", horizontal="center")
            curr_row[5].number_format = '"₹"#,##0.00'
            curr_row[5].alignment = Alignment(vertical="center", horizontal="right")
        curr_row[6].number_format = '"₹"#,##0.00'
        for cell in curr_row:
            cell.border = thin_border

    if txns:
        tot_deposit = sum(float(t.get("amount", 0) or 0) for t in txns if (t.get("type") or "").upper() in ["MONEY_RECEIVED", "DEPOSIT", "CREDIT", "REVERSAL"])
        tot_debited = sum(float(t.get("amount", 0) or 0) for t in txns if (t.get("type") or "").upper() not in ["MONEY_RECEIVED", "DEPOSIT", "CREDIT", "REVERSAL"])
        ws4.append([
            "TOTAL",
            "",
            "",
            "",
            tot_deposit,
            tot_debited,
            "",
            "",
            ""
        ])
        tot_row = ws4[ws4.max_row]
        for c in tot_row:
            c.fill = totals_fill
            c.font = totals_font
            c.border = double_bottom_border
        tot_row[0].alignment = Alignment(vertical="center", horizontal="left")
        tot_row[4].number_format = '"₹"#,##0.00'
        tot_row[4].alignment = Alignment(vertical="center", horizontal="right")
        tot_row[5].number_format = '"₹"#,##0.00'
        tot_row[5].alignment = Alignment(vertical="center", horizontal="right")

    _autofit(ws4)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def _autofit(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = min(len(val_str), 50)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
